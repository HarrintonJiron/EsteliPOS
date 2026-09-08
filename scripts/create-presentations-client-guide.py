from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUTPUT = "/Users/harrintonjiron/agroservicio/Guia rápida - Presentaciones EsteliPOS.docx"
BLUE = "1D4ED8"
NAVY = "172554"
PALE_BLUE = "EFF6FF"
PALE_GREEN = "ECFDF5"
GREEN = "047857"
GRAY = "475569"
LIGHT_GRAY = "E2E8F0"


def font(run, size=11, bold=False, color="0F172A"):
    run.font.name = "Calibri"
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Calibri")
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Calibri")
    run.font.size = Pt(size)
    run.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)


def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def margins(cell, top=100, start=140, bottom=100, end=140):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for side, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_widths(table, widths_dxa):
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths_dxa)))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), "120")
    tbl_ind.set(qn("w:type"), "dxa")
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths_dxa:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for row in table.rows:
        for index, cell in enumerate(row.cells):
            tc_w = cell._tc.get_or_add_tcPr().find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                cell._tc.get_or_add_tcPr().append(tc_w)
            tc_w.set(qn("w:w"), str(widths_dxa[index]))
            tc_w.set(qn("w:type"), "dxa")
            margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def style_paragraph(p, before=0, after=6, line=1.15):
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = line


doc = Document()
section = doc.sections[0]
section.page_width = Inches(8.5)
section.page_height = Inches(11)
section.top_margin = Inches(0.68)
section.bottom_margin = Inches(0.68)
section.left_margin = Inches(0.78)
section.right_margin = Inches(0.78)
section.header_distance = Inches(0.35)
section.footer_distance = Inches(0.35)

normal = doc.styles["Normal"]
normal.font.name = "Calibri"
normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
normal.font.size = Pt(10.5)
normal.paragraph_format.space_after = Pt(5)
normal.paragraph_format.line_spacing = 1.15

for style_name, size, color, before, after in (
    ("Heading 1", 15, BLUE, 12, 6),
    ("Heading 2", 12, NAVY, 8, 4),
):
    style = doc.styles[style_name]
    style.font.name = "Calibri"
    style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    style.font.size = Pt(size)
    style.font.bold = True
    style.font.color.rgb = RGBColor.from_string(color)
    style.paragraph_format.space_before = Pt(before)
    style.paragraph_format.space_after = Pt(after)
    style.paragraph_format.keep_with_next = True

header = section.header.paragraphs[0]
header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
font(header.add_run("ESTELIPOS  |  GUÍA PARA CLIENTES"), 8.5, True, GRAY)

footer = section.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
font(footer.add_run("Northlink Microsystem · Funcionalidad de presentaciones y unidades de medida"), 8, False, GRAY)

kicker = doc.add_paragraph()
style_paragraph(kicker, after=4)
font(kicker.add_run("GUÍA RÁPIDA DE EXPLICACIÓN"), 9, True, BLUE)

title = doc.add_paragraph()
style_paragraph(title, after=2)
font(title.add_run("Ventas por presentación en EsteliPOS"), 23, True, NAVY)

subtitle = doc.add_paragraph()
style_paragraph(subtitle, after=12)
font(subtitle.add_run("Compre en cajas, sacos o quintales; venda en unidades, ristras, libras o la presentación que necesite."), 11.5, False, GRAY)

callout = doc.add_table(rows=1, cols=1)
callout.style = "Table Grid"
set_table_widths(callout, [9360])
shade(callout.cell(0, 0), PALE_GREEN)
p = callout.cell(0, 0).paragraphs[0]
style_paragraph(p, after=0)
font(p.add_run("La idea clave: "), 11, True, GREEN)
font(p.add_run("EsteliPOS guarda el inventario en una unidad base y hace automáticamente las conversiones al comprar o vender."), 11, False, NAVY)

doc.add_heading("Cómo explicárselo al cliente en 30 segundos", level=1)
p = doc.add_paragraph()
style_paragraph(p)
font(p.add_run("“Usted configura una sola vez cómo viene empacado cada producto. Después, el cajero solo elige si vende una unidad, una ristra, una caja, una libra o un quintal. El sistema calcula el precio y descuenta del inventario la cantidad correcta automáticamente.”"), 11, False, NAVY)

doc.add_heading("Ejemplos fáciles", level=1)
table = doc.add_table(rows=1, cols=3)
table.style = "Table Grid"
headers = ["Producto", "Configuración", "Qué hace el sistema"]
for i, text in enumerate(headers):
    shade(table.cell(0, i), BLUE)
    p = table.cell(0, i).paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    style_paragraph(p, after=0)
    font(p.add_run(text), 9.5, True, "FFFFFF")

examples = [
    ("Arroz", "Inventario en libras\n1 quintal = 100 lb", "Compra 2 quintales → agrega 200 lb.\nVende 5 lb → descuenta 5 lb."),
    ("Jabón", "Inventario en unidades\n1 ristra = 3 und\n1 caja = 12 ristras", "Vende 1 caja → descuenta 36 und.\nTambién puede vender por ristra o unidad."),
]
for product, config, result in examples:
    cells = table.add_row().cells
    for i, text in enumerate((product, config, result)):
        if i == 0:
            shade(cells[i], PALE_BLUE)
        p = cells[i].paragraphs[0]
        style_paragraph(p, after=0, line=1.1)
        for idx, line in enumerate(text.split("\n")):
            if idx:
                p.add_run("\n")
            font(p.add_run(line), 9.5, i == 0, NAVY if i == 0 else "0F172A")
set_table_widths(table, [1500, 2850, 5010])

doc.add_heading("Demostración sugerida", level=1)
steps = [
    "Abra el producto y pulse “Compra y venta por presentación”.",
    "Indique cómo controla el inventario: unidad, libra, metro, litro, etc.",
    "Agregue una presentación: por ejemplo, 1 caja contiene 12 ristras.",
    "Marque si se usa para comprar, vender o ambas operaciones y defina su precio.",
    "En el POS seleccione la presentación o escanee su código de barras.",
    "Muestre la comprobación: cuántas unidades base entran o salen del inventario.",
]
for item in steps:
    p = doc.add_paragraph(style="List Number")
    style_paragraph(p, after=3, line=1.12)
    font(p.add_run(item), 10.3)

doc.add_heading("Beneficios que debe destacar", level=1)
benefits = [
    "Un solo producto: no hay que crear “jabón unidad”, “jabón ristra” y “jabón caja” por separado.",
    "Inventario exacto: cada venta descuenta automáticamente su equivalente real.",
    "Rapidez en caja: presentación predeterminada y código de barras por empaque.",
    "Precios flexibles: una caja puede tener descuento sin cambiar el precio unitario.",
    "Historial protegido: cambiar una equivalencia futura no modifica ventas anteriores.",
]
for item in benefits:
    p = doc.add_paragraph(style="List Bullet")
    style_paragraph(p, after=2, line=1.1)
    font(p.add_run(item), 10.2)

note = doc.add_table(rows=1, cols=1)
note.style = "Table Grid"
set_table_widths(note, [9360])
shade(note.cell(0, 0), PALE_BLUE)
p = note.cell(0, 0).paragraphs[0]
style_paragraph(p, after=0)
font(p.add_run("Consejo de demostración: "), 10.5, True, BLUE)
font(p.add_run("use siempre un ejemplo real del negocio del cliente. Para una ferretería: tornillos por unidad, docena y caja; para una miscelánea: arroz por libra y quintal, o jabón por unidad, ristra y caja."), 10.5)

doc.core_properties.title = "Guía rápida - Presentaciones EsteliPOS"
doc.core_properties.subject = "Explicación comercial para clientes"
doc.core_properties.author = "Northlink Microsystem"
doc.save(OUTPUT)
print(OUTPUT)
