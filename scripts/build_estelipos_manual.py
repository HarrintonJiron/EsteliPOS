from pathlib import Path
from datetime import date
from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "docs" / "Manual_de_Usuario_EsteliPOS.docx"
ASSETS = ROOT / "docs" / "manual_assets"
ASSETS.mkdir(parents=True, exist_ok=True)
OUT.parent.mkdir(parents=True, exist_ok=True)

NAVY = "172554"
INDIGO = "4F46E5"
TEAL = "0F766E"
PALE = "EEF2FF"
LIGHT = "F8FAFC"
GRAY = "64748B"
RED = "B91C1C"


def font(size, bold=False):
    candidates = [
        "/System/Library/Fonts/Supplemental/Arial.ttf",
        "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf",
    ]
    if bold:
        candidates = [p.replace("Arial.ttf", "Arial Bold.ttf").replace("DejaVuSans.ttf", "DejaVuSans-Bold.ttf") for p in candidates] + candidates
    for p in candidates:
        if Path(p).exists():
            return ImageFont.truetype(p, size)
    return ImageFont.load_default()


def mock_screen(filename, title, menu, panels, active):
    w, h = 1500, 850
    im = Image.new("RGB", (w, h), "#F1F5F9")
    d = ImageDraw.Draw(im)
    d.rectangle((0, 0, 285, h), fill="#0F172A")
    d.ellipse((28, 25, 92, 89), fill="#FFFFFF")
    d.text((47, 43), "EP", font=font(22, True), fill="#4F46E5")
    d.text((110, 28), "EsteliPOS", font=font(29, True), fill="white")
    d.text((110, 63), "Sistema administrativo", font=font(14), fill="#94A3B8")
    y = 125
    for item in menu:
        if item == active:
            d.rounded_rectangle((18, y-8, 266, y+35), 10, fill="#3730A3")
        d.text((39, y), item, font=font(18, item == active), fill="white" if item == active else "#CBD5E1")
        y += 51
    d.rectangle((285, 0, w, 76), fill="white")
    d.text((325, 22), title, font=font(30, True), fill="#0F172A")
    d.text((1260, 28), "Administrador", font=font(16, True), fill="#334155")
    x, y = 325, 115
    cols = 2
    card_w = 545
    for i, (heading, lines, color) in enumerate(panels):
        cx = x + (i % cols) * 585
        cy = y + (i // cols) * 225
        d.rounded_rectangle((cx, cy, cx+card_w, cy+185), 16, fill="white", outline="#CBD5E1", width=2)
        d.rectangle((cx, cy, cx+8, cy+185), fill=color)
        d.text((cx+30, cy+22), heading, font=font(22, True), fill="#0F172A")
        ly = cy+62
        for line in lines[:5]:
            d.text((cx+34, ly), "• " + line, font=font(16), fill="#475569")
            ly += 27
    im.save(ASSETS / filename, quality=92)


MENU = ["Punto de Venta", "Facturación", "Proformas", "Reparaciones", "Compras", "Inventario", "Clientes", "Proveedores", "Créditos", "Recursos humanos", "Contabilidad", "Reportes", "Configuración"]
mock_screen("mapa_navegacion.png", "Navegación principal", MENU, [
    ("Menú lateral", ["Módulos visibles según permisos", "Opción activa resaltada", "Botón para contraer el menú"], "#4F46E5"),
    ("Barra superior", ["Nombre de pantalla", "Usuario actual", "Accesos rápidos y cierre de sesión"], "#0F766E"),
    ("Área de trabajo", ["Formularios, filtros y resultados", "Mensajes de confirmación o error", "Acciones según el rol"], "#0284C7"),
    ("Regla de seguridad", ["Leer antes de guardar", "No duplicar clics", "Confirmar empresa, bodega y período"], "#B91C1C"),
], "Punto de Venta")
mock_screen("pos.png", "Punto de Venta", MENU, [
    ("Buscar y agregar", ["Nombre o código de barras", "Categorías y tarjetas de producto", "Presentación y bodega de salida"], "#4F46E5"),
    ("Ticket", ["Cantidad, precio y descuento", "Dto. por línea y Quitar", "Cliente y descuento global"], "#0F766E"),
    ("Cobro", ["Efectivo, crédito u otro método", "Monto recibido o referencia", "Confirmar pago una sola vez"], "#0284C7"),
    ("Caja", ["Apartar F4 y recuperar F6", "Corte del día F10", "Cierre Caja mediante arqueo"], "#B45309"),
], "Punto de Venta")
mock_screen("inventario.png", "Inventario y productos", MENU, [
    ("Catálogo", ["Código, nombre y categoría", "Costo, precio, impuesto y estado", "Stock mínimo y fotografía"], "#4F46E5"),
    ("Ubicaciones", ["Bodegas y estantes", "Transferencias entre bodegas", "Disponibilidad por ubicación"], "#0F766E"),
    ("Control", ["Movimientos y ajustes", "Reconciliación física", "Lotes y vencimientos"], "#0284C7"),
    ("Precios y unidades", ["Listas de precios", "Presentaciones y conversiones", "Unidad de venta predeterminada"], "#B45309"),
], "Inventario")
mock_screen("compras.png", "Compras", MENU, [
    ("Encabezado", ["Proveedor y fecha", "Contado o crédito", "Bodega, moneda y tipo de cambio"], "#4F46E5"),
    ("Productos", ["Buscar por nombre o código", "Crear producto rápido", "Cantidad, unidad y costo"], "#0F766E"),
    ("Estados", ["Pendiente o completada", "La recepción afecta existencias", "Revisar antes de cambiar estado"], "#0284C7"),
    ("Documento", ["Guardar compra", "Editar con permiso", "Ticket, PDF o detalle"], "#B45309"),
], "Compras")
mock_screen("clientes_creditos.png", "Clientes y créditos", MENU, [
    ("Cliente", ["Identificación y contacto", "Límite y días de crédito", "Activar crédito solo con autorización"], "#4F46E5"),
    ("Cuenta", ["Saldo y ventas pendientes", "Estado de cuenta", "Facturas vencidas"], "#0F766E"),
    ("Abono", ["Seleccionar cliente", "Registrar importe y referencia", "Emitir comprobante"], "#0284C7"),
    ("Control", ["No aplicar a cliente equivocado", "No superar el saldo", "Conciliar caja y referencia"], "#B91C1C"),
], "Créditos")
mock_screen("rrhh.png", "Recursos humanos y planilla", MENU, [
    ("Empleados", ["Ficha, salario y estado", "Directorio y organigrama", "Turnos y asistencia"], "#4F46E5"),
    ("Movimientos", ["Permisos y ausencias", "Préstamos", "Bonificaciones y deducciones"], "#0F766E"),
    ("Nómina", ["Elegir período", "Revisar ingresos y deducciones", "Generar y pagar con autorización"], "#0284C7"),
    ("Obligaciones", ["INSS", "Aguinaldo", "Evaluaciones de desempeño"], "#B45309"),
], "Recursos humanos")
mock_screen("contabilidad.png", "Contabilidad", MENU, [
    ("Registro", ["Catálogo de cuentas", "Asientos en borrador", "Contabilizar o anular"], "#4F46E5"),
    ("Libros", ["Diario", "Mayor", "Balance de comprobación"], "#0F766E"),
    ("Estados", ["Estado de resultados", "Balance general", "Flujo de caja"], "#0284C7"),
    ("Administración", ["Períodos fiscales", "Centros de costo", "Impuestos y exportaciones"], "#B45309"),
], "Contabilidad")
mock_screen("reportes.png", "Reportes y análisis", MENU, [
    ("Período", ["Hoy, este mes, mes anterior", "Trimestre o fechas manuales", "Generar reporte"], "#4F46E5"),
    ("Filtros", ["Cliente o proveedor", "Estado y condición", "Producto, movimiento y stock"], "#0F766E"),
    ("Lectura", ["Resumen primero", "Detalle después", "Comparar con documentos fuente"], "#0284C7"),
    ("Salida", ["Exportar solo datos revisados", "Conservar período en el nombre", "No confundir ventas con utilidad"], "#B91C1C"),
], "Reportes")
mock_screen("configuracion.png", "Configuración", MENU, [
    ("Negocio", ["Datos generales y apariencia", "Impuestos y tipo de cambio", "Secuencias documentales"], "#4F46E5"),
    ("Acceso", ["Usuarios", "Roles y permisos", "Módulos habilitados"], "#0F766E"),
    ("Seguridad", ["Políticas de acceso", "Contraseñas", "Actividad administrativa"], "#0284C7"),
    ("Zona crítica", ["Restablecimiento del sistema", "Requiere permiso especial", "Puede eliminar información"], "#B91C1C"),
], "Configuración")


doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Inches(8.5), Inches(11)
sec.top_margin, sec.bottom_margin = Inches(0.7), Inches(0.65)
sec.left_margin, sec.right_margin = Inches(0.78), Inches(0.72)

styles = doc.styles
styles["Normal"].font.name = "Aptos"
styles["Normal"].font.size = Pt(10.7)
styles["Normal"].font.color.rgb = RGBColor.from_string("1E293B")
styles["Normal"].paragraph_format.space_after = Pt(6)
styles["Normal"].paragraph_format.line_spacing = 1.12
for name, size in [("Title", 30), ("Heading 1", 22), ("Heading 2", 16), ("Heading 3", 12.5)]:
    s = styles[name]
    s.font.name = "Aptos Display"
    s.font.size = Pt(size)
    s.font.bold = True
    s.font.color.rgb = RGBColor(0, 0, 0)
    s.paragraph_format.space_before = Pt(12)
    s.paragraph_format.space_after = Pt(7)
    s.paragraph_format.keep_with_next = True

if "Caption" in styles:
    styles["Caption"].font.name = "Aptos"
    styles["Caption"].font.size = Pt(9)
    styles["Caption"].font.italic = True
    styles["Caption"].font.color.rgb = RGBColor.from_string(GRAY)


def shade(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tcPr.append(shd)


def cell_margin(cell, top=100, start=110, bottom=100, end=110):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcMar = tcPr.first_child_found_in("w:tcMar")
    if tcMar is None:
        tcMar = OxmlElement("w:tcMar")
        tcPr.append(tcMar)
    for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = OxmlElement("w:" + m)
        node.set(qn("w:w"), str(v)); node.set(qn("w:type"), "dxa")
        tcMar.append(node)


def set_repeat_table_header(row):
    trPr = row._tr.get_or_add_trPr()
    tblHeader = OxmlElement("w:tblHeader")
    tblHeader.set(qn("w:val"), "true")
    trPr.append(tblHeader)


def table(headers, rows, widths=None):
    t = doc.add_table(rows=1, cols=len(headers))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.style = "Table Grid"
    set_repeat_table_header(t.rows[0])
    for i, h in enumerate(headers):
        c = t.rows[0].cells[i]; shade(c, NAVY); cell_margin(c)
        p = c.paragraphs[0]; p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(h); r.bold = True; r.font.color.rgb = RGBColor(255,255,255); r.font.size = Pt(9.2)
    for ri, row in enumerate(rows):
        cells = t.add_row().cells
        for i, val in enumerate(row):
            if ri % 2: shade(cells[i], LIGHT)
            cell_margin(cells[i]); cells[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p = cells[i].paragraphs[0]; p.paragraph_format.space_after = Pt(0)
            r = p.add_run(str(val)); r.font.size = Pt(9.1)
    if widths:
        for row in t.rows:
            for i, width in enumerate(widths): row.cells[i].width = Inches(width)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return t


def bullet(text, level=0):
    p = doc.add_paragraph(style="List Bullet" if level == 0 else "List Bullet 2")
    p.add_run(text)
    return p


def number(text):
    p = doc.add_paragraph(style="List Number")
    p.add_run(text)
    return p


def warning(text):
    p = doc.add_paragraph()
    r = p.add_run("Punto crítico. "); r.bold = True; r.font.color.rgb = RGBColor.from_string(RED)
    p.add_run(text)
    return p


def tip(text):
    p = doc.add_paragraph()
    r = p.add_run("Uso recomendado. "); r.bold = True; r.font.color.rgb = RGBColor.from_string(TEAL)
    p.add_run(text)
    return p


def figure(filename, caption, width=6.8):
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next = True
    p.add_run().add_picture(str(ASSETS / filename), width=Inches(width))
    c = doc.add_paragraph(caption, style="Caption"); c.alignment = WD_ALIGN_PARAGRAPH.CENTER


def page(): doc.add_page_break()


def add_toc():
    topics = [
        "1 Principios de operación segura", "2 Acceso navegación y sesión", "3 Punto de Venta",
        "4 Facturación e historial de ventas", "5 Apertura corte y cierre de caja", "6 Clientes",
        "7 Créditos y abonos", "8 Proformas de venta", "9 Proveedores",
        "10 Compras y recepción de mercadería", "11 Inventario productos y precios",
        "12 Bodegas estantes y transferencias", "13 Movimientos ajustes y reconciliación",
        "14 Reparaciones y gastos de reparación", "15 Recursos humanos y empleados",
        "16 Préstamos bonificaciones deducciones y salarios", "17 Nómina", "18 Contabilidad",
        "19 Reportes y análisis", "20 Configuración y administración",
        "21 Centro de ayuda soporte y cambio de contraseña", "22 Solución de problemas",
        "23 Matriz de errores por mala manipulación", "24 Rutinas recomendadas",
        "25 Lista de verificación por rol", "26 Glosario", "27 Registro de incidencias para soporte",
    ]
    pairs = []
    half = (len(topics) + 1) // 2
    for i in range(half):
        pairs.append([topics[i], topics[i + half] if i + half < len(topics) else ""])
    table(["Capítulos operativos", "Capítulos de control y consulta"], pairs, [3.5, 3.5])


# Cover
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_before = Pt(42)
logo = ROOT / "public" / "images" / "northlink-logo.png"
if logo.exists(): p.add_run().add_picture(str(logo), width=Inches(1.25))
p = doc.add_paragraph("Manual de Usuario EsteliPOS", style="Title"); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p = doc.add_paragraph("Guía completa para la operación segura y correcta del sistema", style="Subtitle"); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_after = Pt(24)
figure("mapa_navegacion.png", "Mapa general de la interfaz de EsteliPOS", 6.4)
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_before = Pt(18)
r = p.add_run("Versión del manual 1.0\nFecha de emisión 21 de septiembre de 2026\nDirigido a cajeros, vendedores, bodegueros, compras, recursos humanos, contabilidad y administración")
r.font.size = Pt(10.5); r.font.color.rgb = RGBColor.from_string(GRAY)
page()

doc.add_heading("Cómo usar este manual", level=1)
doc.add_paragraph("Este manual explica cómo operar EsteliPOS de manera ordenada, segura y verificable. Está escrito para una persona que conoce el negocio, pero que no necesita conocimientos técnicos. El objetivo principal es registrar cada operación una sola vez, en el módulo correcto, con el cliente, proveedor, producto, bodega, período y forma de pago correctos.")
warning("Las opciones visibles dependen de los módulos habilitados y de los permisos del usuario. Si una opción descrita no aparece, no intente acceder mediante enlaces guardados: solicite al administrador que revise su rol.")
table(["Convención", "Significado"], [
    ["Paso numerado", "Acción que debe realizarse en orden."],
    ["Uso recomendado", "Práctica que reduce errores y facilita auditorías."],
    ["Punto crítico", "Operación con impacto en dinero, inventario, nómina, crédito o contabilidad."],
    ["Vista de referencia", "Ilustración basada en la interfaz actual; los datos exactos varían por empresa y permisos."],
], [1.45, 5.55])
doc.add_heading("Contenido", level=2); add_toc()
page()

doc.add_heading("1 Principios de operación segura", level=1)
doc.add_paragraph("Antes de guardar cualquier documento, aplique la regla de las seis verificaciones: usuario, fecha, tercero, ubicación, valores y estado. Una operación correcta debe poder explicarse posteriormente a partir de su documento, referencia y usuario responsable.")
table(["Verificación", "Pregunta antes de guardar", "Riesgo si se omite"], [
    ["Usuario", "¿Estoy usando mi propia cuenta?", "No se puede determinar quién realizó la operación."],
    ["Fecha y período", "¿La fecha pertenece al día o período correcto?", "Reportes y contabilidad quedan en el período equivocado."],
    ["Tercero", "¿Es el cliente, proveedor o empleado correcto?", "Saldos, crédito o historial asignados a otra persona."],
    ["Ubicación", "¿La sucursal, bodega y estante son correctos?", "Existencias disponibles en una ubicación que no corresponde."],
    ["Valores", "¿Cantidad, costo, precio, impuesto, moneda y descuento son correctos?", "Pérdida económica o documentos inconsistentes."],
    ["Estado", "¿Es borrador, pendiente, completado, pagado, contabilizado o anulado?", "Impacto prematuro o duplicado."],
], [1.1, 2.65, 3.25])
doc.add_heading("1.1 Reglas que todo usuario debe seguir", level=2)
for x in [
    "Use una cuenta individual. No comparta contraseña ni PIN.",
    "No use el botón Atrás del navegador durante un guardado o cobro. Espere el mensaje del sistema.",
    "No haga doble clic en Guardar, Confirmar pago, Contabilizar, Pagar o Cerrar caja.",
    "No modifique precios o descuentos para compensar errores de caja o inventario. Corrija la operación con el procedimiento autorizado.",
    "Conserve referencias bancarias, números de factura del proveedor y observaciones suficientes.",
    "Si la sesión vence, recargue e inicie sesión. Verifique primero si el documento se registró antes de repetirlo.",
]: bullet(x)
page()
doc.add_heading("1.2 Diferencia entre ver, crear, editar, eliminar y aprobar", level=2)
table(["Permiso", "Qué permite", "Buenas prácticas"], [
    ["Ver", "Consultar listas, detalles y reportes.", "No asuma que también permite descargar o exportar."],
    ["Crear", "Registrar un documento o catálogo nuevo.", "Busque primero para evitar duplicados."],
    ["Editar", "Cambiar información existente.", "No cambie datos históricos sin comprobar consecuencias."],
    ["Eliminar", "Quitar registros cuando el sistema lo permite.", "Úselo solo para registros incorrectos y sin dependencia."],
    ["Aprobar o pagar", "Confirmar un movimiento con impacto financiero.", "Requiere revisión independiente cuando el negocio lo establezca."],
    ["Exportar", "Descargar datos o reportes.", "Proteja los archivos porque pueden contener información sensible."],
], [1.25, 2.65, 3.1])
page()

doc.add_heading("2 Acceso navegación y sesión", level=1)
doc.add_heading("2.1 Iniciar sesión", level=2)
for s in ["Abra la dirección del sistema proporcionada por la empresa.", "Seleccione su usuario en Selección rápida o escriba su correo o nombre de usuario.", "Elija Con contraseña o Con PIN según su configuración.", "Ingrese la credencial y pulse Entrar a mi negocio una sola vez.", "Confirme que su nombre aparece en el menú lateral."]: number(s)
warning("Si aparece que las credenciales no coinciden, el PIN no está configurado o el usuario está inactivo, no pruebe credenciales de compañeros. Solicite al administrador verificar su usuario.")
doc.add_heading("2.2 Menú principal", level=2)
figure("mapa_navegacion.png", "Vista de referencia 1. Menú lateral, barra superior y área de trabajo")
doc.add_paragraph("El menú se divide en Principal, Operaciones, Gestión y Sistema. Punto de Venta está destacado por ser la función de uso frecuente. El botón lateral permite contraer o ampliar el menú; cuando está contraído, coloque el puntero sobre el icono para ver su nombre.")
page()
doc.add_heading("2.3 Mensajes y botones comunes", level=2)
table(["Elemento", "Función", "Uso correcto"], [
    ["Guardar", "Crea o actualiza un registro.", "Revise campos obligatorios y pulse una vez."],
    ["Cancelar o Regresar", "Sale sin completar la acción actual.", "Úselo antes de guardar si detectó que abrió el formulario equivocado."],
    ["Editar", "Abre un registro existente para modificación.", "Confirme que el código y nombre correspondan."],
    ["Eliminar", "Solicita borrar un registro.", "Lea la confirmación y verifique dependencias."],
    ["Exportar", "Descarga los resultados filtrados.", "Aplique primero fechas y filtros."],
    ["Estado", "Cambia la etapa del documento.", "No avance el estado hasta que la operación física ocurra."],
    ["Imprimir PDF Ticket", "Genera una salida del documento.", "Verifique número, fecha y total antes de entregar."],
], [1.4, 2.25, 3.35])
page()

doc.add_heading("3 Punto de Venta", level=1)
doc.add_paragraph("El Punto de Venta registra ventas rápidas, descuenta existencias de la bodega correspondiente, asigna el documento a un cliente y conserva el método de pago. La secuencia segura es abrir caja, identificar al cliente, agregar productos, revisar el ticket, elegir el pago y confirmar una sola vez.")
figure("pos.png", "Vista de referencia 2. Zonas principales del Punto de Venta")
doc.add_heading("3.1 Antes de vender", level=2)
for x in ["Compruebe que la caja esté abierta y que el fondo inicial sea correcto.", "Confirme la sucursal y, si corresponde, la bodega de salida.", "Revise que el cliente solicite factura a su nombre antes de cobrar.", "Tenga a mano el método de pago y la referencia si no es efectivo."]: bullet(x)
doc.add_heading("3.2 Venta al contado en efectivo", level=2)
for s in ["Busque el producto por nombre o código de barras; pulse Enter para agregarlo.", "Si el producto tiene presentaciones, seleccione la unidad correcta.", "Ajuste la cantidad y verifique precio, descuento e impuesto.", "Seleccione Cliente General o busque al cliente registrado.", "Elija Efectivo y pulse el botón de pago.", "Ingrese Monto recibido o use Exacto o los botones de billetes.", "Revise total y cambio. Pulse Confirmar Pago una sola vez.", "Entregue el recibo o ticket y conserve el efectivo en caja."]: number(s)
doc.add_heading("3.3 Venta por transferencia u otro método", level=2)
doc.add_paragraph("Seleccione el método correspondiente. Registre el número de referencia exactamente como aparece en el comprobante y use Notas para identificar banco, últimos dígitos o situación especial según la política de la empresa. No marque una transferencia como recibida basándose únicamente en una captura no verificada.")
doc.add_heading("3.4 Venta a crédito", level=2)
for s in ["Seleccione un cliente registrado; Cliente General no debe usarse para crédito.", "Confirme que el crédito esté habilitado y revise límite, saldo disponible y días autorizados.", "Agregue productos y seleccione Crédito como condición o método.", "Explique al cliente la fecha de vencimiento y entregue su comprobante.", "Si el límite se excede, use Solicitar autorización solamente con el administrador presente. La opción Autorizar una vez no modifica permanentemente el límite."]: number(s)
warning("Nunca use la autorización de otro administrador sin su intervención. Un exceso de crédito aumenta la exposición de cobro y debe quedar justificado.")
doc.add_heading("3.5 Descuentos", level=2)
table(["Opción", "Alcance", "Control recomendado"], [
    ["Dto. en línea", "Solo el producto seleccionado.", "Verifique el porcentaje entre 0 y 100 y la política comercial."],
    ["Descuento de factura", "Todo el ticket, porcentual o fijo.", "Compruebe el total final y evite duplicar un descuento ya incluido en el precio."],
    ["Quitar descuento", "Retira el descuento aplicado.", "Úselo antes de cobrar; después debe seguirse el procedimiento de corrección."],
], [1.55, 2.15, 3.3])
doc.add_heading("3.6 Apartar recuperar y descartar tickets", level=2)
doc.add_paragraph("Apartar F4 conserva temporalmente un ticket sin convertirlo en venta. Recuperar F6 permite continuar uno apartado. Verifique cliente y productos al recuperarlo. Eliminar un ticket apartado o Descartar limpia el trabajo actual; no equivale a anular una venta ya confirmada.")
page()
doc.add_heading("3.7 Botones y atajos del POS", level=2)
table(["Botón u opción", "Función", "Precaución"], [
    ["F1 Atajos", "Muestra ayuda de teclas rápidas.", "Úselo para aprender sin modificar el ticket."],
    ["Apartar F4", "Guarda temporalmente el ticket.", "No entrega mercancía ni registra ingreso."],
    ["Recuperar F6", "Reabre un ticket apartado.", "No confunda dos clientes."],
    ["Descuento", "Abre descuento por factura o producto.", "Requiere autorización comercial cuando aplique."],
    ["Descartar", "Vacía el ticket actual.", "Confirme que no estaba listo para cobro."],
    ["Corte del Día F10", "Consulta actividad de ventas del día.", "Es informativo; no sustituye el arqueo."],
    ["Cierre Caja", "Abre el flujo de arqueo.", "Cuente efectivo antes de confirmar."],
], [1.55, 2.35, 3.1])
page()

doc.add_heading("4 Facturación e historial de ventas", level=1)
doc.add_paragraph("Facturación muestra los documentos ya registrados. Se usa para buscar, consultar detalle, imprimir recibo, PDF o ticket, y editar o eliminar cuando el rol y el estado lo permiten.")
for s in ["Abra Facturación desde Operaciones.", "Use fecha, cliente, número o estado para localizar la venta.", "Abra Ver o el número del documento y confirme sus detalles.", "Use Recibo, Imprimir o PDF según el formato requerido.", "Edite solo cuando la política y el estado lo permitan; documente el motivo."]: number(s)
warning("No elimine una venta para corregir una diferencia de caja sin confirmar el efecto en inventario, crédito y contabilidad. Si ya se entregó el producto o se emitió documento fiscal, siga el procedimiento administrativo de anulación o devolución definido por la empresa.")
doc.add_heading("4.1 Cambio y comprobantes", level=2)
doc.add_paragraph("La pantalla de cambio permite verificar cuánto debe entregarse al cliente después del cobro. El recibo térmico está diseñado para 80 mm y debe imprimirse sin depender de Internet. Antes de reimprimir, confirme que corresponde al número de venta correcto.")

doc.add_heading("5 Apertura corte y cierre de caja", level=1)
doc.add_heading("5.1 Apertura", level=2)
for s in ["Cuente el fondo inicial físicamente.", "Abra Arqueo o use la opción relacionada desde el POS.", "Registre el importe exacto y confirme una sola vez.", "No mezcle dinero personal ni caja de otro turno."]: number(s)
doc.add_heading("5.2 Corte del día", level=2)
doc.add_paragraph("El corte resume ventas por método de pago y sirve para revisión durante el turno. No cierra la sesión de caja. Compare efectivo esperado, transferencias y créditos con los documentos generados.")
doc.add_heading("5.3 Cierre y arqueo", level=2)
for s in ["Detenga ventas temporalmente y cuente efectivo por denominación.", "Verifique vouchers, referencias y salidas autorizadas.", "Ingrese el monto contado, no el monto esperado por el sistema.", "Revise diferencia sobrante o faltante y escriba una explicación verificable.", "Confirme Cierre de Caja una sola vez y conserve el reporte."]: number(s)
warning("No cambie ventas, descuentos o métodos de pago para forzar una diferencia a cero. La diferencia debe investigarse y documentarse.")
page()

doc.add_heading("6 Clientes", level=1)
doc.add_paragraph("El catálogo de clientes centraliza identificación, contacto y condiciones de crédito. Antes de crear uno, búsquelo por nombre, cédula, RUC, teléfono o correo para evitar duplicados.")
table(["Campo", "Uso correcto"], [
    ["Nombre", "Nombre de la persona o nombre comercial identificable."],
    ["Tipo", "Natural o empresa, según corresponda."],
    ["Cédula RUC", "Documento sin inventar valores; respete el formato empresarial."],
    ["Nombre de empresa", "Razón o denominación usada en documentos."],
    ["Teléfono correo dirección", "Datos para contacto y cobro."],
    ["Crédito habilitado", "Actívelo solo después de aprobación."],
    ["Límite y días", "Capacidad máxima y plazo pactado."],
], [2.05, 5.0])
doc.add_heading("6.1 Crear editar y consultar", level=2)
for s in ["Abra Clientes y busque primero.", "Pulse Crear cliente o use Cliente Rápido desde POS si solo necesita los datos esenciales.", "Complete identificación y contacto; revise errores de digitación.", "Guarde y abra el detalle para confirmar.", "Para corregir, pulse Editar; no cree otro registro."]: number(s)
warning("Cambiar la habilitación o límite de crédito afecta ventas futuras. No se debe usar para ocultar una deuda existente ni para permitir una venta sin autorización.")

doc.add_heading("7 Créditos y abonos", level=1)
figure("clientes_creditos.png", "Vista de referencia 3. Relación entre cliente, saldo, abono y control")
doc.add_heading("7.1 Consultar deuda", level=2)
doc.add_paragraph("Use Créditos para ver saldos por cliente, buscar cuentas, abrir el detalle, consultar Estado de cuenta y revisar Vencidos. El reporte permite filtrar y, con permiso, exportar.")
doc.add_heading("7.2 Registrar un abono", level=2)
for s in ["Localice al cliente desde Créditos.", "Abra Nuevo abono.", "Confirme saldo y documentos pendientes.", "Ingrese monto, fecha, método y referencia.", "Revise que el abono no se aplique al cliente equivocado.", "Guarde una vez y emita el comprobante de pago."]: number(s)
warning("No registre un abono si el dinero o transferencia no está confirmado. No utilice un abono negativo ni otro cliente para compensar errores; escale la corrección.")
page()

doc.add_heading("8 Proformas de venta", level=1)
doc.add_paragraph("Una proforma comunica una cotización y no debe tratarse como venta pagada ni como salida de inventario. Permite crear, consultar, editar, eliminar según permisos, convertir a venta y emitir ticket o PDF.")
for s in ["Abra Proformas y seleccione Nueva proforma.", "Identifique al cliente y vigencia de la oferta.", "Agregue productos, cantidades, precios y descuentos autorizados.", "Revise observaciones, impuestos y total.", "Guarde y entregue el ticket o PDF marcado como proforma.", "Cuando el cliente acepte, use Convertir a venta y vuelva a revisar existencias y pago."]: number(s)
warning("No entregue mercancía únicamente con una proforma. La salida debe quedar respaldada por una venta confirmada.")

doc.add_heading("9 Proveedores", level=1)
doc.add_paragraph("El proveedor identifica a quien abastece el negocio y conserva sus condiciones de pago. El producto se vincula operativamente con el proveedor al registrarlo en una compra; crear un proveedor no aumenta inventario por sí solo.")
table(["Grupo", "Campos y criterio"], [
    ["Identificación", "Código, nombre comercial, razón social y RUC."],
    ["Contacto", "Persona, teléfono, correo, ciudad y dirección."],
    ["Condiciones", "Tipo, contado o crédito de 15, 30 o 60 días, límite y estado."],
], [1.65, 5.35])
doc.add_heading("9.1 Crear proveedor", level=2)
for s in ["Abra Proveedores y busque por nombre o RUC.", "Pulse Crear proveedor.", "Complete al menos nombre comercial y datos necesarios para compras.", "Elija la condición real de pago y el estado Activo.", "Guarde y abra el detalle para confirmar."]: number(s)
doc.add_heading("9.2 Editar desactivar o eliminar", level=2)
doc.add_paragraph("Use Editar para actualizar contacto o condiciones. Prefiera Inactivo cuando el proveedor tiene historial y ya no se utiliza. Eliminar debe reservarse para duplicados o registros sin dependencias; el sistema puede impedirlo si existen compras relacionadas.")
doc.add_heading("9.3 Crear productos vinculados al proveedor", level=2)
doc.add_paragraph("La forma más directa es iniciar una compra, seleccionar el proveedor y usar Nuevo producto o Crear desde búsqueda. Complete código, nombre, costo de compra, precio de venta, unidad y categoría. Al guardar la compra, el historial del producto quedará relacionado con ese proveedor por el documento de compra.")
warning("No cree el mismo producto una vez por cada proveedor. Un producto debe conservar un código único; los distintos abastecimientos se registran mediante compras.")
page()

doc.add_heading("10 Compras y recepción de mercadería", level=1)
figure("compras.png", "Vista de referencia 4. Datos que deben revisarse en una compra")
doc.add_heading("10.1 Registrar compra", level=2)
for s in ["Abra Compras y pulse Nueva compra.", "Seleccione proveedor; use Proveedor rápido solo si confirmó que no existe.", "Ingrese fecha y elija pago a crédito, contado en efectivo o contado por transferencia.", "Seleccione bodega destino.", "Revise moneda y tipo de cambio cuando difiera de la moneda de la empresa.", "Busque productos; agregue cantidad, unidad y costo documentado.", "Si falta un producto, use Nuevo producto y complete su información mínima.", "Compare subtotal, impuestos y total contra la factura del proveedor.", "Guarde y cambie el estado únicamente cuando la mercadería haya sido recibida según el flujo del negocio."]: number(s)
doc.add_heading("10.2 Compra a crédito", level=2)
doc.add_paragraph("Seleccione A crédito por pagar. Verifique la condición acordada con el proveedor y conserve el número de factura. El documento pendiente debe aparecer en el control correspondiente; no lo marque como pagado solo porque la mercadería llegó.")
doc.add_heading("10.3 Estados edición y eliminación", level=2)
doc.add_paragraph("Los cambios de estado controlan la etapa de la compra. Antes de editar o eliminar una compra completada, confirme si ya actualizó inventario y contabilidad. Una corrección mal aplicada puede duplicar entradas o dejar el stock sin respaldo.")
warning("La bodega destino define dónde queda disponible la mercadería. Una bodega incorrecta produce faltantes aparentes en una ubicación y sobrantes en otra.")

doc.add_heading("11 Inventario productos y precios", level=1)
figure("inventario.png", "Vista de referencia 5. Catálogo, ubicaciones, control y precios")
doc.add_heading("11.1 Crear producto en modo completo", level=2)
table(["Sección", "Qué registrar", "Error frecuente"], [
    ["Información básica", "Código único, nombre, categoría, unidad y descripción.", "Duplicar el código o usar nombres ambiguos."],
    ["Imagen", "Foto clara del producto real.", "Usar imagen de otra presentación."],
    ["Precios y stock", "Costo, precio, stock mínimo, estado e impuesto.", "Confundir costo con precio de venta."],
    ["Promoción", "Descuento y etiqueta vigentes.", "Dejar promoción activa fuera de fecha."],
    ["Agroquímicos", "Lote, vencimiento, registro sanitario, ingrediente y concentración.", "Omitir trazabilidad o fecha de vencimiento."],
    ["Observaciones", "Información operativa útil.", "Registrar instrucciones no verificadas."],
], [1.55, 3.1, 2.3])
for s in ["Abra Inventario y pulse Nuevo producto o Producto rápido.", "Use Siguiente código cuando esté disponible o ingrese el código oficial.", "Seleccione unidad base; no la cambie después de movimientos sin análisis.", "Registre costo y precio de venta. Revise impuesto y margen.", "Defina stock mínimo y estado Activo.", "Guarde; si necesita presentaciones, use Guardar y configurar presentaciones."]: number(s)
doc.add_heading("11.2 Precio de compra y precio de venta", level=2)
doc.add_paragraph("El precio de compra representa el costo unitario de adquisición. El precio de venta es el importe ofrecido al cliente antes o después de impuestos según la configuración. No copie automáticamente el costo como precio. Revise margen, impuesto, flete y política comercial.")
doc.add_heading("11.3 Editar producto", level=2)
doc.add_paragraph("Busque el producto y pulse Editar. Los cambios de nombre, imagen, promoción o stock mínimo suelen ser administrativos. Cambios de código, unidad base, costo, impuesto o precio deben revisarse porque afectan búsqueda, presentaciones, márgenes y ventas futuras.")
doc.add_heading("11.4 Unidades presentaciones y conversiones", level=2)
doc.add_paragraph("Unidades define catálogos como unidad, caja o kilogramo. Las conversiones permiten vender una presentación equivalente a varias unidades base. Configure el factor con cuidado y marque una unidad de venta predeterminada. Convertir unidades modifica la composición del stock; registre el motivo y verifique existencias antes y después.")
warning("Un factor invertido, por ejemplo 1 unidad igual a 12 cajas en lugar de 1 caja igual a 12 unidades, puede multiplicar o reducir el inventario de forma grave.")
doc.add_heading("11.5 Listas de precios", level=2)
doc.add_paragraph("Cree una lista con nombre, vigencia y criterio comercial. Abra su detalle y agregue productos con sus precios. Al editar o quitar un ítem, compruebe qué clientes o ventas utilizan la lista. Mantenga una sola fuente vigente para cada política comercial.")
doc.add_heading("11.6 Carga masiva y exportación", level=2)
doc.add_paragraph("La carga masiva acelera el alta de muchos productos, pero debe validarse en una muestra antes de confirmar. Revise códigos únicos, formatos numéricos, categorías y unidades. Exportar permite auditoría o análisis; el archivo no sustituye la información del sistema.")
page()

doc.add_heading("12 Bodegas estantes y transferencias", level=1)
doc.add_heading("12.1 Bodegas", level=2)
doc.add_paragraph("Una bodega representa una ubicación de existencias. Cree nombre, código y datos de ubicación claros. La bodega principal debe reflejar el almacén operativo. No elimine una bodega con stock o movimientos; transfiera y verifique primero.")
doc.add_heading("12.2 Estantes", level=2)
doc.add_paragraph("Los estantes ayudan a localizar físicamente el producto dentro de una bodega. Use códigos visibles en el almacén y en el sistema. Editar un nombre debe acompañarse de señalización física para evitar ubicaciones desactualizadas.")
doc.add_heading("12.3 Transferencia entre bodegas", level=2)
for s in ["Abra Transferencias y seleccione origen y destino distintos.", "Consulte disponibilidad del producto en la bodega de origen.", "Ingrese cantidad y unidad correctas.", "Registre referencia u observación de traslado.", "Confirme cuando el movimiento físico esté autorizado.", "En destino, verifique recepción y cantidades."]: number(s)
warning("Una transferencia no es compra ni venta: reduce una ubicación y aumenta otra sin cambiar el total global. No la use para corregir faltantes sin investigación.")

doc.add_heading("13 Movimientos ajustes y reconciliación", level=1)
doc.add_heading("13.1 Movimientos", level=2)
doc.add_paragraph("Movimientos es el historial de entradas y salidas generadas por compras, ventas, ajustes, conversiones o transferencias. Filtre por producto, tipo, fecha y ubicación para investigar diferencias.")
doc.add_heading("13.2 Ajustes de inventario", level=2)
for s in ["Realice conteo físico independiente.", "Abra Ajustes y seleccione el producto exacto.", "Elija entrada o salida según la diferencia.", "Ingrese cantidad y motivo verificable: daño, vencimiento, conteo u otro autorizado.", "Guarde una vez y revise el movimiento generado."]: number(s)
warning("Nunca use un ajuste para registrar una compra o venta omitida. Corrija el documento de origen cuando sea posible, porque el ajuste no conserva la misma información comercial o contable.")
doc.add_heading("13.3 Reconciliación", level=2)
doc.add_paragraph("Reconciliar alinea existencias con un conteo controlado. Debe programarse, limitar movimientos durante el conteo y conservar evidencia. Compare cantidades antes de confirmar, especialmente productos de alto valor, agroquímicos, lotes y vencimientos.")
page()

doc.add_heading("14 Reparaciones y gastos de reparación", level=1)
doc.add_paragraph("Reparaciones controla órdenes de servicio, equipo recibido, diagnóstico, servicios, repuestos, estados, gastos y entrega al cliente.")
for s in ["Abra Reparaciones y pulse Nueva reparación.", "Seleccione o cree el cliente correcto.", "Registre equipo, marca, identificación, accesorios, condición y falla informada.", "Agregue servicios y artículos con cantidades y precios.", "Guarde y entregue el ticket de recepción.", "Actualice el estado conforme avance el trabajo, no por anticipado.", "Registre gastos relacionados con comprobante y categoría.", "Antes de entregar, revise total, pagos, estado y conformidad del cliente."]: number(s)
table(["Estado o acción", "Uso recomendado"], [
    ["Recibido o pendiente", "El equipo está bajo custodia y aún no ha terminado el diagnóstico."],
    ["En proceso", "Hay trabajo autorizado en curso."],
    ["Completado", "El servicio terminó y fue revisado."],
    ["Entregado", "El cliente recibió el equipo; registre fecha y respaldo."],
    ["Ticket PDF", "Comprobante para cliente o archivo."],
    ["Gastos", "Costos operativos asociados; no confundir con cobro al cliente."],
], [2.05, 4.95])
warning("No cambie una reparación a Entregado si el equipo continúa físicamente en el taller. La custodia debe coincidir con el estado del sistema.")

doc.add_heading("15 Recursos humanos y empleados", level=1)
figure("rrhh.png", "Vista de referencia 6. Componentes de recursos humanos y planilla")
doc.add_heading("15.1 Ficha del empleado", level=2)
doc.add_paragraph("Cree una ficha por empleado con datos personales, cargo, área, fecha de ingreso, salario y estado. Evite duplicados por cambio de cargo: edite la ficha existente y conserve los soportes del cambio.")
doc.add_heading("15.2 Directorio organigrama turnos y asistencia", level=2)
doc.add_paragraph("El Directorio facilita contacto; el Organigrama muestra relaciones de puestos; Turnos presenta horarios; Asistencia registra presencia o incidencias. Seleccione empleado y fecha correctos antes de guardar asistencia. No edite asistencia para cuadrar una nómina sin autorización y evidencia.")
doc.add_heading("15.3 Permisos y ausencias", level=2)
for s in ["Abra Permisos y cree una solicitud.", "Seleccione empleado, tipo, fechas y motivo.", "Adjunte o conserve respaldo según la política interna.", "Revise el estado antes de calcular nómina.", "Edite o elimine solo mientras el proceso y permisos lo permitan."]: number(s)
doc.add_heading("15.4 Evaluaciones INSS y aguinaldo", level=2)
doc.add_paragraph("Evaluaciones registra resultados de desempeño. INSS y Aguinaldo muestran cálculos o referencias del período según los datos disponibles. Revise la normativa aplicable y los parámetros vigentes antes de tratar un cálculo como definitivo; el sistema ayuda al proceso, pero no reemplaza la revisión responsable de nómina.")
page()

doc.add_heading("16 Préstamos bonificaciones deducciones y salarios", level=1)
table(["Movimiento", "Qué representa", "Control clave"], [
    ["Salario", "Base remunerativa del empleado.", "Fecha de vigencia y autorización."],
    ["Préstamo", "Monto entregado y saldo por descontar.", "Cuota, plazo y comprobante."],
    ["Bonificación", "Ingreso adicional.", "Motivo, período, aprobación y pago."],
    ["Deducción", "Reducción autorizada.", "Base legal o autorización, monto y período."],
], [1.45, 2.65, 2.9])
doc.add_heading("16.1 Crear y procesar movimientos", level=2)
for s in ["Seleccione siempre el empleado desde el catálogo.", "Indique concepto, monto, fecha o período y observación.", "Revise si el movimiento es recurrente o de una sola vez.", "Use Aprobar únicamente después de comprobar el respaldo.", "Use Marcar pagado cuando el desembolso se haya realizado.", "Verifique el impacto en la nómina del período."]: number(s)
warning("No reutilice un movimiento de otro empleado ni cambie el monto después de pagado sin seguir un procedimiento de corrección. Los movimientos de nómina contienen información sensible.")

doc.add_heading("17 Nómina", level=1)
doc.add_paragraph("La nómina reúne salario, ingresos adicionales y deducciones para un período. La pantalla permite elegir fechas, ver comparativo por empleado, desglose de deducciones, tendencia reciente, detalle y ticket de nómina.")
for s in ["Abra Recursos humanos y luego Nómina.", "Seleccione el período y pulse Ver período.", "Revise empleados activos y salario base.", "Verifique asistencia, permisos, préstamos, bonificaciones y deducciones.", "Compare total bruto, deducciones y neto por empleado.", "Investigue valores atípicos antes de aprobar o pagar.", "Genere el ticket o comprobante correspondiente."]: number(s)
warning("No procese nómina con un período equivocado ni con empleados duplicados o inactivos. Un cambio de salario debe registrarse y autorizarse antes del cálculo.")
doc.add_heading("17.1 Errores frecuentes", level=2)
for x in ["Bonificación duplicada por guardar dos veces.", "Cuota de préstamo aplicada después de liquidar el saldo.", "Deducción asignada al empleado equivocado.", "Ausencia sin registrar que aumenta el pago.", "Cambio de salario sin fecha de vigencia clara.", "Marcar como pagada una bonificación todavía pendiente."]: bullet(x)
page()

doc.add_heading("18 Contabilidad", level=1)
figure("contabilidad.png", "Vista de referencia 7. Registro, libros, estados y administración contable")
doc.add_paragraph("Contabilidad transforma y consulta los efectos financieros de la operación. Las ventas no son lo mismo que la utilidad; la utilidad considera costo de ventas y gastos. La caja disponible tampoco equivale a ganancia.")
doc.add_heading("18.1 Catálogo de cuentas", level=2)
doc.add_paragraph("Cuentas organiza activos, pasivos, patrimonio, ingresos y gastos. Cree o edite cuentas solo con criterio contable. No elimine cuentas con movimientos; desactívelas o siga la política establecida.")
doc.add_heading("18.2 Asientos", level=2)
for s in ["Abra Asientos y pulse Nuevo asiento.", "Seleccione fecha dentro de un período abierto.", "Escriba una descripción verificable y referencia.", "Agregue líneas con cuenta, debe, haber y centro de costo cuando corresponda.", "Compruebe que total Debe sea igual a total Haber.", "Guarde como borrador para revisión.", "Use Contabilizar solo cuando esté aprobado. Use Anular en lugar de borrar un asiento ya contabilizado cuando corresponda."]: number(s)
warning("Contabilizar fija el asiento en los libros. No lo haga para probar el sistema. Un asiento descuadrado, duplicado o en período incorrecto distorsiona todos los estados financieros.")
doc.add_heading("18.3 Diario mayor y balance de comprobación", level=2)
doc.add_paragraph("Diario muestra los asientos en orden cronológico. Mayor agrupa movimientos por cuenta. Balance de comprobación resume débitos, créditos y saldos. Use el mismo rango de fechas para comparar y exporte únicamente después de revisar filtros.")
doc.add_heading("18.4 Estado de resultados", level=2)
doc.add_paragraph("Muestra ingresos, costos, gastos y resultado del período. Para consultar día, semana o mes, seleccione las fechas correspondientes. Una pérdida puede deberse a costos o gastos elevados, período incompleto, documentos sin contabilizar o clasificación incorrecta; investigue antes de concluir.")
doc.add_heading("18.5 Balance general y flujo de caja", level=2)
doc.add_paragraph("Balance general presenta activos, pasivos y patrimonio a una fecha. Flujo de caja ayuda a entender entradas y salidas de efectivo. Un negocio puede mostrar utilidad y poco efectivo si vende a crédito o mantiene inventario alto.")
doc.add_heading("18.6 Períodos centros de costo e impuestos", level=2)
doc.add_paragraph("Períodos fiscales delimitan dónde se registran los asientos. Cerrar un período evita cambios posteriores y requiere revisión. Centros de costo permiten analizar áreas, sucursales o actividades. Impuestos deben configurarse con tasa, vigencia y tratamiento correctos; no cambie una tasa histórica para corregir documentos anteriores.")
page()

doc.add_heading("19 Reportes y análisis", level=1)
figure("reportes.png", "Vista de referencia 8. Períodos, filtros, lectura y exportación")
doc.add_heading("19.1 Consultar día semana o mes", level=2)
for s in ["Abra Reportes.", "Elija la pestaña del reporte: ventas, compras, inventario u otra disponible.", "Use Hoy, Este mes, Mes anterior o Trimestre, o escriba Desde y Hasta.", "Aplique filtros por cliente, proveedor, estado, condición, producto, movimiento o stock según el reporte.", "Pulse Generar reporte.", "Lea primero el resumen y luego el detalle.", "Exporte solo si los filtros son correctos."]: number(s)
table(["Pregunta", "Reporte o módulo recomendado", "Cuidado"], [
    ["¿Cuánto vendimos?", "Ventas o Facturación por período.", "El total de ventas no es utilidad."],
    ["¿Cuánto ganamos o perdimos?", "Estado de resultados.", "Requiere costos y gastos correctamente registrados."],
    ["¿Qué se movió hoy?", "Movimientos de inventario y Diario contable.", "Diferencie entrada, salida y transferencia."],
    ["¿Qué falta en bodega?", "Inventario con Stock bajo o Sin stock.", "Confirme conteo físico."],
    ["¿Qué vence pronto?", "Inventario Por vencer o Vencido.", "Revise lote y fecha real."],
    ["¿Quién debe?", "Créditos, Vencidos y Estado de cuenta.", "Considere abonos aún no conciliados."],
], [2.05, 2.7, 2.25])
doc.add_heading("19.2 Dashboard analítica y sucursales", level=2)
doc.add_paragraph("Dashboard ofrece indicadores operativos; Analítica gerencial ayuda a observar tendencias y comparaciones; Sucursales separa información por establecimiento cuando está configurado. Verifique siempre el período, la sucursal y si el indicador usa ventas brutas, netas, cobros o utilidad.")
page()

doc.add_heading("20 Configuración y administración", level=1)
figure("configuracion.png", "Vista de referencia 9. Áreas administrativas y zona crítica")
doc.add_heading("20.1 Datos generales apariencia e imágenes", level=2)
doc.add_paragraph("General contiene identidad y parámetros del negocio. Apariencia controla tema, color, nombre y logotipo. Use archivos claros y autorizados. Un cambio visual no modifica movimientos, pero puede afectar la identificación de tickets y pantallas.")
doc.add_heading("20.2 Tipos de cambio e impuestos", level=2)
doc.add_paragraph("Registre la fecha y tasa correctas. Evite borrar tasas utilizadas en compras. En Impuestos, configure nombre, porcentaje, vigencia y modo de visualización. Verifique el tratamiento antes de cambiar valores, pues afectará documentos nuevos.")
doc.add_heading("20.3 Secuencias", level=2)
doc.add_paragraph("Secuencias controla la numeración de documentos. No cambie el siguiente número para reutilizar un consecutivo ni para ocultar una anulación. Los saltos deben investigarse y documentarse.")
doc.add_heading("20.4 Módulos", level=2)
doc.add_paragraph("Módulos activa o desactiva áreas funcionales. Desactivar un módulo oculta su uso, pero no debe interpretarse como eliminación de datos. Revise dependencias y usuarios antes de guardar.")
doc.add_heading("20.5 Usuarios roles y permisos", level=2)
for s in ["Cree una cuenta individual con nombre y correo correctos.", "Asigne el rol mínimo necesario.", "Compare roles antes de clonar o ampliar privilegios.", "Use Activar o desactivar para altas y bajas temporales.", "Restablezca la contraseña solo por un proceso autorizado; el usuario debe cambiarla cuando corresponda.", "Revise periódicamente permisos de crear, editar, eliminar, aprobar, pagar y exportar."]: number(s)
warning("No asigne Administrador para resolver rápidamente un acceso. Conceda el permiso concreto y vuelva a comprobar el menú del usuario.")
doc.add_heading("20.6 Seguridad y restablecimiento", level=2)
doc.add_paragraph("Seguridad reúne controles administrativos. Restablecimiento del sistema es una opción excepcional con permiso especial y limitación de intentos. Puede afectar información de negocio y no forma parte del trabajo diario.")
warning("No ejecute Restablecimiento del sistema sin respaldo verificado, autorización escrita, alcance definido y ventana de mantenimiento. Puede eliminar o reiniciar datos.")
page()

doc.add_heading("21 Centro de ayuda soporte y cambio de contraseña", level=1)
doc.add_paragraph("Centro de ayuda concentra orientación y contacto de soporte. Antes de solicitar ayuda, anote módulo, acción, fecha y hora, número de documento, mensaje exacto y usuario. No envíe contraseñas ni bases de datos completas por medios no autorizados.")
doc.add_heading("21.1 Cambio de contraseña", level=2)
doc.add_paragraph("Use la pantalla Cambiar contraseña cuando el sistema lo requiera o cuando sospeche exposición. Elija una contraseña exclusiva y guárdela de forma segura. Cierre sesiones abiertas en equipos compartidos.")

doc.add_heading("22 Solución de problemas", level=1)
table(["Situación", "Causa probable", "Qué hacer"], [
    ["Opción no visible", "Módulo desactivado o permiso insuficiente.", "Solicitar revisión del rol; no usar enlaces directos."],
    ["Sesión vencida mensaje 419", "Inactividad o reinicio.", "Iniciar sesión y comprobar si el documento ya existe antes de repetir."],
    ["No hay stock", "Bodega incorrecta, producto agotado o movimiento pendiente.", "Consultar disponibilidad y movimientos; no inventar cantidad."],
    ["Código duplicado", "El producto ya existe.", "Buscarlo y editarlo; no cambiar el código para duplicar."],
    ["No se puede eliminar", "Registro relacionado con documentos.", "Desactivar o corregir mediante el flujo autorizado."],
    ["Crédito rechazado", "No habilitado, vencido o sin saldo disponible.", "Revisar cliente; solicitar autorización solo si procede."],
    ["Compra no aumenta stock", "Estado pendiente o bodega distinta.", "Revisar estado y destino antes de crear otra compra."],
    ["Caja no cuadra", "Método incorrecto, vuelto, omisión o conteo.", "Recontar, comparar documentos y registrar diferencia."],
    ["Reporte vacío", "Fechas, filtros o sucursal sin datos.", "Limpiar filtros y validar período."],
    ["Asiento no contabiliza", "Descuadre, período cerrado o permiso.", "Revisar debe y haber, fecha y autorización."],
], [1.75, 2.25, 3.0])
doc.add_heading("22.1 Protocolo ante doble clic o pantalla congelada", level=2)
for s in ["No vuelva a pulsar Guardar o Confirmar inmediatamente.", "Espere la respuesta y revise mensajes.", "Abra la lista del módulo en otra navegación segura y busque por número, fecha, tercero y total.", "Si el documento existe, no lo repita.", "Si no existe y no hubo impacto físico, vuelva a registrar con precaución.", "Si hay duda sobre dinero, inventario o contabilidad, detenga el proceso y escale."]: number(s)
page()

doc.add_heading("23 Matriz de errores por mala manipulación", level=1)
table(["Mala práctica", "Consecuencia", "Prevención"], [
    ["Compartir usuarios", "Auditoría sin responsable confiable.", "Cuenta individual y bloqueo al retirarse."],
    ["Crear duplicados", "Saldos e historiales fragmentados.", "Buscar antes de crear."],
    ["Vender desde bodega equivocada", "Stock falso por ubicación.", "Confirmar bodega antes de cobrar."],
    ["Confundir costo y precio", "Pérdida o margen distorsionado.", "Revisar ambos campos y el impuesto."],
    ["Aplicar doble descuento", "Venta por debajo de lo autorizado.", "Revisar línea y descuento global."],
    ["Repetir pago por sesión vencida", "Venta o abono duplicado.", "Buscar el documento antes de reintentar."],
    ["Completar compra antes de recibir", "Existencia disponible que no está físicamente.", "Cambiar estado al recibir."],
    ["Ajustar para cuadrar", "Oculta causa de faltante o sobrante.", "Investigar y documentar."],
    ["Cerrar caja con monto esperado", "Oculta diferencia real.", "Ingresar conteo físico."],
    ["Contabilizar sin revisión", "Estados financieros incorrectos.", "Borrador, revisión y aprobación."],
    ["Procesar nómina con datos incompletos", "Pago incorrecto y reclamos.", "Cierre de novedades antes de cálculo."],
    ["Cambiar secuencias", "Duplicidad o pérdida de trazabilidad.", "Solo administrador autorizado."],
], [2.15, 2.45, 2.45])

doc.add_heading("24 Rutinas recomendadas", level=1)
doc.add_heading("24.1 Inicio del día", level=2)
for x in ["Iniciar sesión con usuario propio.", "Revisar fecha, sucursal y conectividad.", "Abrir caja con fondo contado.", "Confirmar impresora y papel si se usan tickets.", "Revisar alertas de stock, vencimientos y pendientes críticos."]: bullet(x)
doc.add_heading("24.2 Durante el día", level=2)
for x in ["Registrar en tiempo real; no acumular ventas, compras o abonos en papel.", "Usar referencias y observaciones suficientes.", "Mantener documentos físicos ordenados por tipo y fecha.", "Investigar mensajes antes de repetir una acción.", "No dejar sesión abierta sin supervisión."]: bullet(x)
doc.add_heading("24.3 Cierre diario", level=2)
for x in ["Terminar ventas pendientes y recuperar o eliminar tickets apartados según corresponda.", "Ejecutar corte y comparar métodos de pago.", "Contar efectivo y cerrar caja con diferencia real.", "Revisar ventas, abonos, compras completadas y ajustes del día.", "Entregar soportes y cerrar sesión."]: bullet(x)
doc.add_heading("24.4 Revisión semanal y mensual", level=2)
table(["Frecuencia", "Revisión"], [
    ["Semanal", "Créditos vencidos, stock bajo, movimientos inusuales, compras pendientes, reparaciones demoradas, asistencia y novedades de nómina."],
    ["Mensual", "Estado de resultados, balance, flujo de caja, conciliaciones, inventario sensible, nómina, impuestos, usuarios y permisos."],
], [1.25, 5.75])
page()

doc.add_heading("25 Lista de verificación por rol", level=1)
table(["Rol", "Responsabilidades mínimas"], [
    ["Cajero", "Abrir caja, identificar cliente, revisar ticket, cobrar una vez, emitir recibo, corte y arqueo."],
    ["Vendedor", "Producto y presentación correctos, descuento autorizado, cliente y condición de pago correctos."],
    ["Bodega", "Recepción, ubicación, transferencias, conteos, lotes, vencimientos y movimientos documentados."],
    ["Compras", "Proveedor, factura, condición, moneda, costos, bodega y estado de recepción."],
    ["Cobranza", "Cliente correcto, saldo, abono confirmado, referencia, comprobante y vencidos."],
    ["Recursos humanos", "Ficha única, asistencia, permisos, salario, novedades y confidencialidad."],
    ["Planilla", "Período, ingresos, deducciones, préstamos, aprobación, pago y comprobantes."],
    ["Contabilidad", "Períodos, cuentas, asientos balanceados, contabilización, anulaciones, estados y exportaciones."],
    ["Administrador", "Usuarios, roles, módulos, impuestos, tasas, secuencias, respaldos y cambios críticos."],
], [1.45, 5.55])

doc.add_heading("26 Glosario", level=1)
table(["Término", "Definición sencilla"], [
    ["Arqueo", "Comparación entre el dinero contado y el esperado por el sistema."],
    ["Bodega", "Ubicación donde se controlan existencias."],
    ["Centro de costo", "Clasificación para analizar gastos o resultados por área."],
    ["Contabilizar", "Confirmar un asiento para que forme parte de los libros."],
    ["Costo", "Valor de adquisición o producción del producto."],
    ["Crédito", "Venta o compra cuyo pago queda pendiente."],
    ["Debe y Haber", "Dos lados de un asiento que deben mantener equilibrio."],
    ["Existencia Stock", "Cantidad disponible de un producto."],
    ["Lista de precios", "Conjunto de precios para una política o grupo comercial."],
    ["Proforma", "Cotización que todavía no es una venta confirmada."],
    ["Reconciliación", "Comparación y ajuste controlado entre conteo físico y sistema."],
    ["Secuencia", "Numeración consecutiva usada por los documentos."],
], [1.7, 5.3])

doc.add_heading("27 Registro de incidencias para soporte", level=1)
doc.add_paragraph("Copie esta información cuando solicite ayuda. No incluya contraseñas ni PIN.")
table(["Dato", "Espacio para completar"], [
    ["Fecha y hora", ""], ["Usuario", ""], ["Sucursal y bodega", ""], ["Módulo y pantalla", ""],
    ["Acción realizada", ""], ["Número de documento", ""], ["Mensaje exacto", ""],
    ["Resultado esperado", ""], ["Resultado observado", ""], ["¿Hubo impacto en dinero o stock?", ""],
], [2.5, 4.5])
doc.add_paragraph("Fin del manual. Este documento debe actualizarse cuando cambien pantallas, reglas de negocio, impuestos, permisos o flujos críticos.")

# Headers, footers and document settings
for section in doc.sections:
    hp = section.header.paragraphs[0]
    hp.text = "EsteliPOS  |  Manual de usuario"
    hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    hp.runs[0].font.size = Pt(8.5); hp.runs[0].font.color.rgb = RGBColor.from_string(GRAY)
    fp = section.footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = fp.add_run("Northlink Microsystem  •  EsteliPOS  •  Página ")
    r.font.size = Pt(8.5); r.font.color.rgb = RGBColor.from_string(GRAY)
    fld = OxmlElement("w:fldSimple"); fld.set(qn("w:instr"), "PAGE")
    fp._p.append(fld)

settings = doc.settings._element
update = OxmlElement("w:updateFields"); update.set(qn("w:val"), "true"); settings.append(update)

# Set core properties
doc.core_properties.title = "Manual de Usuario EsteliPOS"
doc.core_properties.subject = "Guía completa para la operación segura y correcta de EsteliPOS"
doc.core_properties.author = "Northlink Microsystem"
doc.core_properties.keywords = "EsteliPOS, manual de usuario, punto de venta, inventario, compras, planilla, contabilidad"
doc.core_properties.comments = "Elaborado a partir de los módulos, rutas y pantallas disponibles en la aplicación EsteliPOS."

doc.save(OUT)
print(OUT)
